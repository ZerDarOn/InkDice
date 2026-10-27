"""Read-only extraction. Never execute workbook formulas or import a player's card."""
import json, hashlib, re, sys, zipfile, xml.etree.ElementTree as ET
from pathlib import Path
import openpyxl

source = Path(sys.argv[1])
target = Path(__file__).resolve().parents[1] / 'prototype/coc/catalog.mjs'
w = openpyxl.load_workbook(source, data_only=False)
def text(value):
    return '' if value is None else str(value).strip()
def cell(sheet, row, col):
    return text(w[sheet].cell(row, col).value)
def entry(kind, name, sheet, area, **fields):
    key=f'{sheet}:{area}' if kind in ['reference','era','experience'] else area
    return dict(id=f'coc7:{kind}:{key}', kind=kind, name=name, sourceSheet=sheet, sourceRange=area, fields=fields)

entries, occupations, skills, weapons, armors = [], [], [], [], []
matrix=w['本职技能']
for r in range(4,233):
    number=w['职业列表'].cell(r,1).value
    if not isinstance(number,int): continue
    col=next((c for c in range(2,257) if matrix.cell(1,c).value==number),None)
    marks={text(matrix.cell(i,1).value):text(matrix.cell(i,col).value) for i in range(3,73) if col and matrix.cell(i,col).value is not None}
    formula=cell('职业列表',r,6)
    assert re.fullmatch(r'=[A-Z0-9*+(),]+',formula), formula
    e=entry('occupation',cell('职业列表',r,2),'职业列表',f'A{r}:M{r}',信用范围=cell('职业列表',r,4),职业属性=cell('职业列表',r,5),本职技能=cell('职业列表',r,7),推荐关系人=cell('职业列表',r,11),介绍=cell('职业列表',r,13))
    e.update(number=number,formula=formula,marks=marks)
    occupations.append(e); entries.append(e)
for r in range(95,212):
    name=cell('附表',r,2)
    if not name or name.startswith('=') or '——' in name: continue
    headers=['初始值','特殊适用','难度依据','说明','普通难度','困难难度','孤注一掷示例','孤注一掷失败','疯狂时']
    e=entry('skill',name,'附表',f'B{r}:K{r}',**{h:cell('附表',r,c) for c,h in enumerate(headers,3)})
    skills.append(e);entries.append(e)
for r in range(2,106):
    name=cell('武器列表 战斗',r,21)
    if not name:continue
    keys=['技能','伤害','射程','贯穿','每轮','装弹量','故障值','常见时代','价格20s/现代','发明时间']
    e=entry('weapon',name,'武器列表 战斗',f'B{r}:U{r}',**{k:cell('武器列表 战斗',r,c) for c,k in enumerate(keys,3)})
    weapons.append(e);entries.append(e)
group=''
for r in range(2,76):
    name=cell('防具表 载具表',r,2)
    if not name:continue
    group=cell('防具表 载具表',r,10).replace('\n','') or group
    e=entry('armor',name,'防具表 载具表',f'B{r}:J{r}',**{k:cell('防具表 载具表',r,c) for c,k in enumerate(['护甲值','MOV惩罚','覆盖位置','使用物种','防利器穿刺','常见时代','价格20s/现代'],3)},分类=group)
    armors.append(e);entries.append(e)
# Reference sections preserve prose and original cell addresses, not computed values
# tied to this specific investigator. No identity, portrait, inventory or campaign logs.
reference_sections=[('技能注释','reference',1,154,1,21),('属性注释','reference',1,192,2,45),('资产及物价参考','price',1,229,1,29),('武器列表 战斗','combat',107,339,1,37),('防具表 载具表','vehicle',77,355,1,34),('疯狂表','sanity',1,145,1,8),('附表','experience',19,24,2,5),('附表','reference',39,48,4,5)]
for sheet,kind,r1,r2,c1,c2 in reference_sections:
    for row in w[sheet].iter_rows(min_row=r1,max_row=r2,min_col=c1,max_col=c2):
        values=[c for c in row if c.value is not None and c.data_type!='f']
        if not values:continue
        strings=[text(c.value) for c in values if isinstance(c.value,str)]
        if not strings:continue
        name=next((s for s in strings if len(s)>1 and not re.fullmatch(r'[\d\W]+',s)),strings[0])[:65]
        area=f'{values[0].coordinate}:{values[-1].coordinate}'
        entries.append(entry(kind,name,sheet,area,原表内容='\n'.join(f'{c.coordinate}　{text(c.value)}' for c in values)))
# Core sheet skill rows: defaults are read from the template, never from allocated points.
core=[]
for col,base_col,spec_col in [(6,10,8),(28,32,30)]:
    for r in range(16,50):
        label=cell('人物卡',r,col)
        if not label:continue
        default=w['人物卡'].cell(r,base_col).value
        specialization=bool(re.search(r'[①②③：]',label)) or label in ['母语','技艺①','技艺②','技艺③']
        base=default if isinstance(default,(int,float)) else (0 if label=='自定义技能' else 25 if label=='格斗：' else 20 if label=='射击：' else 5 if '技艺' in label else 1)
        core.append(dict(id=f'skill-{col}-{r}',label=label,base=base,specialization=specialization,baseRule='dexHalf' if label=='闪避' else 'edu' if label=='母语' else '',defaultSpecialization='斗殴' if label=='格斗：' else '手枪' if label=='射击：' else ''))
# Read literal validation lists, without relying on openpyxl's unsupported x14 extension.
option_cells={'H20:I20':'技艺','H35:I36':'格斗','H39:I40':'射击','AD31:AE32':'科学','AD27:AE27':'驾驶'}
skill_options={}
with zipfile.ZipFile(source) as z:
    xml=ET.fromstring(z.read('xl/worksheets/sheet3.xml'))
    for element in xml.iter():
        if element.tag.split('}')[-1]!='dataValidation': continue
        group=option_cells.get(element.attrib.get('sqref'))
        formula=next((e.text for e in element.iter() if e.tag.split('}')[-1]=='formula1'),None)
        if group and formula and formula.startswith('"'):
            skill_options[group]=[s for s in formula.strip('"').split(',') if s]
skill_options['格斗'].insert(0,'斗殴');skill_options['射击'].insert(0,'手枪')
# Year rows contain source reference facts. Omit cumulative invention lists AQ/AR.
for r in range(2,492):
    year=w['附表'].cell(r,31).value
    if not isinstance(year,(int,float)): continue
    cells=[w['附表'].cell(r,c) for c in list(range(32,43))+[45]]
    values=[c for c in cells if c.value is not None and c.data_type!='f']
    if values: entries.append(entry('era',f'{year}年','附表',f'AE{r}:AS{r}',原表内容='\n'.join(f'{c.coordinate}　{text(c.value)}' for c in values)))
data=dict(ruleset='coc-7e',source='CoC七版 CY23PlusFinal（用户提供工作簿）',sha256=hashlib.sha256(source.read_bytes()).hexdigest(),notice='原表资料含作者扩充及可选规则，未逐条确认为官方；保留原表来源。未导入原人物的私人内容。',occupations=occupations,skills=skills,weapons=weapons,armors=armors,coreSkills=core,skillOptions=skill_options,entries=entries)
target.parent.mkdir(parents=True,exist_ok=True)
target.write_text('// Generated by tools/import-coc-workbook.py. Do not edit by hand.\nexport const COC_CATALOG = '+json.dumps(data,ensure_ascii=False,separators=(',',':'))+';\n',encoding='utf-8')
print(json.dumps({k:len(data[k]) for k in ['occupations','skills','weapons','armors','coreSkills','entries']},ensure_ascii=False))
