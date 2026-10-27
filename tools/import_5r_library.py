"""Read-only import of the supplied 2024 character workbook into the local library.

Every populated source cell is retained with its address. Only columns whose
meaning is clear from the workbook headers are exposed as searchable entries.
No formulas are evaluated and the source workbook is never saved.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
import warnings
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from import_5e_library import ROOT, DEFAULT_DB, as_text, json_value


ROLES = {
    "主要": "character_template", "起源": "character_template", "背包": "character_template",
    "盟友与魔宠": "character_template", "法术书": "character_template", "导入": "utility",
    "据点": "character_template", "自定义调整栏": "utility", "更新": "log",
    "法术大全": "reference", "职业": "reference", "装备": "reference", "种族": "reference",
    "专长与据点": "reference", "魔宠": "reference", "背景": "reference", "数据表": "calculator",
    "圆形效应范围": "utility",
}
SKIP_NAMES = {"", "-", "—", "0", "无", "请选择", "物品", "装备名称", "法术名", "背景", "种族"}
DICE = re.compile(r"^(?:\d+d\d+(?:\s*/\s*\d+d\d+)?|\d+)$", re.IGNORECASE)


def name_ok(value):
    return isinstance(value, str) and value.strip() not in SKIP_NAMES and not value.strip().startswith(("—", "自定义武器"))


def add_entry(db, source_id, kind, name, sheet, cell_range, payload, rank=None, english=None, book=None):
    entry_id = f"{source_id}:{sheet}:{cell_range}:{kind}"
    db.execute(
        """INSERT INTO library_entries
        (id,source_id,ruleset,kind,name,english_name,rank,source_book,source_sheet,source_range,payload_json)
        VALUES (?,?,'dnd5r-2024',?,?,?,?,?,?,?,?)
        ON CONFLICT(id) DO UPDATE SET name=excluded.name, payload_json=excluded.payload_json
        WHERE library_entries.name IS NOT excluded.name
           OR library_entries.payload_json IS NOT excluded.payload_json""",
        (entry_id, source_id, kind, name.strip(), english, rank, book, sheet, cell_range, json_value(payload)),
    )


def import_spells(db, source_id, sheet):
    count = 0
    headers = [as_text(cell.value) for cell in sheet[2]]
    for number, row in enumerate(sheet.iter_rows(min_row=3, values_only=False), 3):
        values = [cell.value for cell in row]
        name, rank = values[0], values[1]
        if not name_ok(name) or not isinstance(rank, int) or not 0 <= rank <= 9:
            continue
        payload = {header: values[index] for index, header in enumerate(headers) if header}
        payload["职业列表"] = [headers[index] for index in range(14, min(25, len(headers)))
                             if as_text(values[index]).lower() not in ("", "×", "x", "0", "-")]
        add_entry(db, source_id, "spell", name, sheet.title, f"A{number}:AA{number}",
                  payload, rank, as_text(values[13]) or None, as_text(values[25]) or None)
        count += 1
    return count


def import_equipment(db, source_id, sheet):
    counts = Counter()
    armor_category = None
    for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=False), 2):
        values = [cell.value for cell in row]
        # AH:AO: the workbook's actual weapon catalogue, including 2024 mastery.
        weapon, damage = values[33], values[34]
        if name_ok(weapon) and DICE.fullmatch(as_text(damage)):
            add_entry(db, source_id, "weapon", weapon, sheet.title, f"AH{number}:AO{number}", {
                "伤害骰": damage, "伤害类型": values[35], "特性": values[36],
                "精通": values[37], "重量": values[38], "价格": values[39], "使用属性": values[40],
            })
            counts["weapon"] += 1
        # A:D: adventuring gear. Other lookup blocks are intentionally not folded in.
        item = values[0]
        if name_ok(item) and not as_text(item).startswith(("——", "—")) and any(values[i] is not None for i in (1, 2, 3)):
            add_entry(db, source_id, "item", item, sheet.title, f"A{number}:D{number}", {
                "描述": values[1], "价格": values[2], "重量": values[3],
            })
            counts["item"] += 1
        # L:Q is a tool catalogue, separate from adventuring gear.
        tool = values[11]
        if name_ok(tool) and as_text(values[12]) and number <= 30:
            add_entry(db, source_id, "tool", tool, sheet.title, f"L{number}:Q{number}", {
                "关联属性": values[12], "操作": values[13], "制造": values[14],
                "价格": values[15], "重量": values[16],
            })
            counts["tool"] += 1
        # Z:AA names weapon masteries and explains their effects.
        mastery = values[25]
        if name_ok(mastery) and isinstance(values[26], str) and len(values[26].strip()) > 15:
            add_entry(db, source_id, "weapon_mastery", mastery, sheet.title,
                      f"Z{number}:AA{number}", {"效果": values[26]})
            counts["weapon_mastery"] += 1
        # AP:AT mixes actual armor with natural defenses and formulas. Only
        # explicit light/medium/heavy equipment rows are catalogued as armor.
        if number in (28, 33, 40):
            armor_category = {28: "轻甲", 33: "中甲", 40: "重甲"}[number]
        armor = values[41]
        if armor_category and 29 <= number <= 45 and name_ok(armor) and isinstance(values[42], (int, float)):
            add_entry(db, source_id, "armor", armor, sheet.title, f"AP{number}:AT{number}", {
                "类型": armor_category, "基础AC": values[42],
                "敏捷加值公式格": f"AR{number}", "重量": values[44],
                "特性": values[45], "说明": "原表动态公式另存于单元格快照，未移植为护甲计算规则。",
            })
            counts["armor"] += 1
    return counts


def import_backgrounds(db, source_id, sheet):
    count = 0
    for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=False), 2):
        values = [cell.value for cell in row]
        # I:O is a VLOOKUP result for the currently selected character, not a
        # second catalogue. Only B:H contains independent background records.
        name = values[1]
        if not name_ok(name) or not as_text(values[2]):
            continue
        add_entry(db, source_id, "background", name, sheet.title,
                  f"B{number}:H{number}", {
                      "属性值": values[2], "专长": values[3],
                      "技能熟练": values[4], "工具熟练": values[5],
                      "装备": values[6], "背景描述": values[7],
                  })
        count += 1
    return count


def import_species(db, source_id, sheet):
    count = 0
    stats = {as_text(row[4]): (number, row[5:8])
             for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2)
             if as_text(row[4])}
    family = ""
    for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=False), 2):
        values = [cell.value for cell in row]
        if not name_ok(values[0]):
            continue
        if isinstance(values[1], int):
            family = as_text(sheet.cell(values[1] * 2, 11).value)
        if not family:
            raise ValueError(f"Missing species family at A{number}")
        stat_row, common = stats.get(family, (None, (None, None, None)))
        add_entry(db, source_id, "species", values[0], sheet.title, f"A{number}:H{number}", {
            "通用特性": family, "生物种类": common[0], "体型": common[1], "速度": common[2],
            "族系": family, "选项层级": "基础种族" if values[0] == family else "亚种/血统选项",
            "数值来源": f"E{stat_row}:H{stat_row}" if stat_row else "原表无同名通用数值",
            "说明": "数值为族系通用信息；亚种覆盖效果由玩家确认。",
        })
        count += 1
    return count


def import_species_traits(db, source_id, sheet):
    """Name/description pairs in the K:AV matrix; choices remain player-owned."""
    count = 0
    rows = list(sheet.iter_rows(values_only=True))
    base_names = {as_text(row[4]) for row in rows[1:]}
    for number in range(2, len(rows), 2):
        species = as_text(rows[number - 1][10])
        if species not in base_names or not species:
            continue
        for column in range(12, 49):
            name = rows[number - 1][column - 1]
            description = rows[number][column - 1]
            if not name_ok(name) or not isinstance(description, str) or len(description.strip()) < 8:
                continue
            letter = get_column_letter(column)
            add_entry(db, source_id, "species_trait", name, sheet.title,
                      f"{letter}{number}:{letter}{number+1}", {
                          "种族": species, "描述": description,
                          "说明": "来源表的特性名称/文本；子种选择与等级条件未自动应用。",
                      })
            count += 1
    return count


def import_classes(db, source_id, sheet):
    count = 0
    rows = list(sheet.iter_rows(values_only=True))
    hit_dice = {as_text(row[40]): as_text(row[41]) for row in rows[1:21] if as_text(row[40])}
    for offset, first_col in enumerate(range(1, 56, 3)):
        number = offset + 3
        row = rows[number - 1]
        name = rows[32][first_col]
        if not name_ok(name):
            continue
        subclasses = [as_text(value) for value in row[2:26] if name_ok(value)]
        proficiency = {as_text(rows[index][first_col + 1]): rows[index][first_col + 2]
                       for index in range(33, 39) if as_text(rows[index][first_col + 1])}
        add_entry(db, source_id, "class", name, sheet.title, f"B{number}:Z{number}", {
            "子职业名称": subclasses, "生命骰": hit_dice.get(as_text(name)),
            "初始熟练与装备": proficiency,
            "扩展职业开关": number >= 15,
            "子职名称来源": f"C{number}:Z{number}",
            "初始熟练与装备来源": f"{get_column_letter(first_col + 1)}33:"
                               f"{get_column_letter(first_col + 3)}39",
        })
        count += 1
    return count


def import_class_features(db, source_id, sheet):
    count = 0
    rows = list(sheet.iter_rows(values_only=True))
    for first_col in range(1, 56, 3):
        class_name = rows[32][first_col]
        if not name_ok(class_name):
            continue
        for number in range(40, min(285, len(rows)) + 1):
            row = rows[number - 1]
            level, name, description = row[first_col:first_col + 3]
            if not name_ok(name) or not isinstance(description, str) or len(description.strip()) < 8:
                continue
            valid_level = level if isinstance(level, int) and 1 <= level <= 20 else None
            add_entry(db, source_id, "class_feature", name, sheet.title,
                      f"{get_column_letter(first_col + 1)}{number}:"
                      f"{get_column_letter(first_col + 3)}{number}", {
                "职业": class_name, "表内等级": valid_level, "描述": description,
                "说明": "含子职/扩展能力；仅供检索，不能按等级自动授予。",
            })
            count += 1
    return count


def import_spell_progressions(db, source_id, sheet):
    """Four explicit 1–20 progression grids, retained as lookups, not rules."""
    count = 0
    tables = (
        ("全施法者", "BE", "BN", "BO"),
        ("魔契师", "BP", "BU", None),
        ("渎魂圣约", "BV", "BZ", None),
        ("药剂师", "CA", "CF", None),
    )
    rows = list(sheet.iter_rows(values_only=True))
    for mode, start, end, points_col in tables:
        first = column_index_from_string(start) - 1
        last = column_index_from_string(end) - 1
        for number in range(3, 23):
            row = rows[number - 1]
            level = row[first]
            if not isinstance(level, int) or not 1 <= level <= 20:
                continue
            slots = {str(rank): row[first + rank] for rank in range(1, last - first + 1)
                     if isinstance(row[first + rank], int)}
            points = row[column_index_from_string(points_col) - 1] if points_col else None
            cell_end = points_col or end
            add_entry(db, source_id, "spell_progression", f"{mode}·{level}级", sheet.title,
                      f"{start}{number}:{cell_end}{number}", {
                "模式": mode, "等级": level, "法术位": slots,
                "法术值": points if isinstance(points, int) else None,
                "说明": "来源表静态进度；多职业、恢复和例外尚未自动计算。",
            })
            count += 1
    return count


def import_feats(db, source_id, sheet):
    count = 0
    for number, row in enumerate(sheet.iter_rows(min_row=4, values_only=False), 4):
        for name_index, description_index, category in ((1, 2, "专长"), (3, 4, "战斗风格"), (5, 6, "专长"), (7, 8, "专长")):
            name, description = row[name_index].value, row[description_index].value
            if not name_ok(name) or not isinstance(description, str) or len(description.strip()) < 8:
                continue
            first = chr(65 + name_index)
            last = chr(65 + description_index)
            add_entry(db, source_id, "feat", name, sheet.title, f"{first}{number}:{last}{number}", {
                "分类": category, "描述": description,
            })
            count += 1
    return count


def import_bastion_facilities(db, source_id, sheet):
    count = 0
    for number, row in enumerate(sheet.iter_rows(min_row=7, values_only=False), 7):
        name, prerequisite, space, staff, order, description = [row[index].value for index in range(10, 16)]
        if not name_ok(name) or not isinstance(description, str) or len(description.strip()) < 20:
            continue
        add_entry(db, source_id, "bastion", name, sheet.title, f"K{number}:P{number}", {
            "先决条件": prerequisite, "空间": space, "雇员": staff,
            "指令": order, "描述": description,
        })
        count += 1
    return count


def import_familiars(db, source_id, sheet):
    count = 0
    for number, row in enumerate(sheet.iter_rows(min_row=3, values_only=False), 3):
        values = [cell.value for cell in row]
        name = values[0]
        if not name_ok(name) or not as_text(values[1]) or not isinstance(values[4], (int, float)):
            continue
        add_entry(db, source_id, "creature", name, sheet.title, f"A{number}:AD{number}", {
            "体型": values[1], "类型": values[2], "阵营": values[3], "AC": values[4],
            "HP": values[5], "HD": values[6], "先攻": values[7], "速度": values[8],
            "六维": {key: values[index] for key, index in (("str", 9), ("dex", 12), ("con", 15),
                                                        ("int", 18), ("wis", 21), ("cha", 24))},
            "概览": values[27], "特质": values[28], "动作": values[29],
        })
        count += 1
    return count


IMPORTERS = {
    "法术大全": lambda db, source, sheet: Counter(spell=import_spells(db, source, sheet)),
    "装备": import_equipment,
    "背景": lambda db, source, sheet: Counter(background=import_backgrounds(db, source, sheet)),
    "种族": lambda db, source, sheet: Counter(species=import_species(db, source, sheet),
                                             species_trait=import_species_traits(db, source, sheet)),
    "职业": lambda db, source, sheet: Counter(class_=import_classes(db, source, sheet),
                                            class_feature=import_class_features(db, source, sheet),
                                            spell_progression=import_spell_progressions(db, source, sheet)),
    "专长与据点": lambda db, source, sheet: Counter(feat=import_feats(db, source, sheet),
                                                  bastion=import_bastion_facilities(db, source, sheet)),
    "魔宠": lambda db, source, sheet: Counter(creature=import_familiars(db, source, sheet)),
}


def import_workbook(source: Path, destination: Path = DEFAULT_DB):
    source = source.resolve(strict=True)
    if source.suffix.lower() != ".xlsx":
        raise ValueError("5R 来源必须是 .xlsx 工作簿")
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    source_id = f"5r-workbook-{digest[:16]}"
    destination.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(destination)
    db.execute("PRAGMA foreign_keys = ON")
    db.executescript((ROOT / "schemas" / "content-library.sql").read_text(encoding="utf-8"))
    existing = db.execute("SELECT id FROM library_sources WHERE sha256 = ?", (digest,)).fetchone()
    if existing:
        before_changes = db.total_changes
        cached = load_workbook(source, read_only=True, data_only=True, keep_links=False)
        try:
            with db:
                db.execute("DELETE FROM library_entries WHERE source_id=? AND source_sheet='背景' "
                           "AND source_range='I2:O2' AND kind='background'", (existing[0],))
                for sheet_name in ("装备", "种族", "专长与据点", "背景", "职业"):
                    IMPORTERS[sheet_name](db, existing[0], cached[sheet_name])
        finally:
            cached.close()
        counts = dict(db.execute("SELECT kind,COUNT(*) FROM library_entries WHERE source_id=? GROUP BY kind", (existing[0],)))
        cells = db.execute("SELECT COUNT(*) FROM library_cells WHERE source_id=?", (existing[0],)).fetchone()[0]
        changed = db.total_changes > before_changes
        db.close()
        return {"status": "updated" if changed else "already_imported", "source_id": existing[0], "cells": cells, "entries": counts}

    warnings.simplefilter("ignore", UserWarning)
    cached = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    formulas = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    counts = Counter()
    cell_count = 0
    try:
        with db:
            db.execute("INSERT INTO library_sources VALUES (?,?,?,?,?,?,?)", (
                source_id, source.name, digest, "dnd5r-2024", datetime.now(timezone.utc).isoformat(),
                "local_only", "用户提供的 2024 人物卡；含扩展内容，未经逐条规则或分发授权复核。",
            ))
            for sheet in cached.worksheets:
                db.execute("INSERT INTO library_sheets VALUES (?,?,?,?,?,?)", (
                    source_id, sheet.title, ROLES.get(sheet.title, "unclassified"),
                    sheet.sheet_state, sheet.max_row, sheet.max_column,
                ))
                batch = []
                for value_row, formula_row in zip(sheet.iter_rows(), formulas[sheet.title].iter_rows()):
                    for value_cell, formula_cell in zip(value_row, formula_row):
                        value = value_cell.value
                        formula = formula_cell.value if formula_cell.data_type == "f" else None
                        if value is None and formula is None:
                            continue
                        if formula is not None and not isinstance(formula, str):
                            formula = json_value({"text": getattr(formula, "text", None),
                                                  "ref": getattr(formula, "ref", None)})
                        batch.append((source_id, sheet.title, value_cell.coordinate,
                                      json_value(value) if value is not None else None, formula))
                        if len(batch) >= 1000:
                            db.executemany("INSERT INTO library_cells VALUES (?,?,?,?,?)", batch)
                            cell_count += len(batch)
                            batch.clear()
                if batch:
                    db.executemany("INSERT INTO library_cells VALUES (?,?,?,?,?)", batch)
                    cell_count += len(batch)
                importer = IMPORTERS.get(sheet.title)
                if importer:
                    counts.update(importer(db, source_id, sheet))
    finally:
        cached.close()
        formulas.close()
        db.close()
    # The class importer uses a Python-safe keyword in Counter construction.
    if "class_" in counts:
        counts["class"] = counts.pop("class_")
    return {"status": "imported", "source_id": source_id, "cells": cell_count,
            "sheets": len(cached.sheetnames), "entries": dict(counts), "database": str(destination)}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="5R 原始工作簿（只读）")
    parser.add_argument("--database", type=Path, default=DEFAULT_DB)
    args = parser.parse_args()
    print(json.dumps(import_workbook(args.source, args.database), ensure_ascii=False))
