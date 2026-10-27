"""Read-only import of the 5E workbook into a local, provenance-preserving library.

The workbook is never saved, its VBA is never loaded, and external links are ignored.
Every nonempty worksheet cell is archived. Only unambiguous tables become entries.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
import warnings
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DB = ROOT / "local-data" / "content-library.sqlite"
SHEET_ROLES = {
    "法术大全": "reference",
    "图表": "reference",
    "属性计算器&自定义调整栏": "calculator",
    "角色": "character_template",
    "主要情况": "character_template",
    "背包": "character_template",
    "部署生物": "character_template",
    "施法": "character_template",
    "录卡": "character_template",
    "导入海豹骰或梨骰": "utility",
    "更新日志&在线更新": "log",
}
PLACEHOLDERS = {"", "-", "—", "0", "未填写", "请选择", "请先选择物品"}


def as_text(value):
    return "" if value is None else str(value).strip()


def is_name(value):
    return isinstance(value, str) and bool(value.strip()) and value.strip() not in PLACEHOLDERS


def json_value(value):
    if isinstance(value, (datetime, date)):
        value = value.isoformat()
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), default=str)


def add_entry(cursor, source_id, kind, name, row, sheet, end_col, payload,
              english_name=None, rank=None, source_book=None):
    entry_id = f"{source_id}:{sheet}:{row}:{kind}"
    start_col = {"spell": "A", "class": "C", "weapon": "U", "item": "AB",
                 "background": "AI", "feat": "CY"}[kind]
    cursor.execute(
        """INSERT INTO library_entries
        (id, source_id, ruleset, kind, name, english_name, rank, source_book,
         source_sheet, source_range, payload_json)
        VALUES (?, ?, 'dnd5e-2014', ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(id) DO UPDATE SET
          source_range=excluded.source_range, payload_json=excluded.payload_json
        WHERE library_entries.source_range IS NOT excluded.source_range
           OR library_entries.payload_json IS NOT excluded.payload_json""",
        (entry_id, source_id, kind, name, english_name, rank, source_book,
         sheet, f"{start_col}{row}:{end_col}{row}", json_value(payload)),
    )


def add_cell_entry(cursor, source_id, kind, name, sheet, cell_range, payload):
    cursor.execute(
        """INSERT INTO library_entries
        (id, source_id, ruleset, kind, name, source_sheet, source_range, payload_json)
        VALUES (?, ?, 'dnd5e-2014', ?, ?, ?, ?, ?)
        ON CONFLICT(id) DO NOTHING""",
        (f"{source_id}:{sheet}:{cell_range}:{kind}", source_id, kind,
         name.strip(), sheet, cell_range, json_value(payload)),
    )


def import_spells(cursor, source_id, sheet):
    headers = [cell.value for cell in sheet[1]]
    count = 0
    for row in sheet.iter_rows(min_row=2, values_only=False):
        cells = [cell.value for cell in row]
        name, rank = cells[0], cells[1]
        if not is_name(name) or not isinstance(rank, int) or not 0 <= rank <= 9:
            continue  # excludes blank reserve rows and the ten "列表0–9" headings
        fields = {str(header): cells[index] for index, header in enumerate(headers) if header}
        classes = [str(headers[index]) for index in range(13, 24)
                   if as_text(cells[index]) not in ("", "×", "x", "X", "0", "-")]
        fields["职业列表"] = classes
        add_entry(cursor, source_id, "spell", name.strip(), row[0].row,
                  "法术大全", "Y", fields, as_text(cells[12]) or None,
                  rank, as_text(cells[24]) or None)
        count += 1
    return count


def import_chart_entries(cursor, source_id, sheet):
    counts = Counter()
    rows = list(sheet.iter_rows(values_only=True))
    # C1:E13 is the hit-die/first-level HP table. C14: onward is a different table.
    for index in range(13):
        row = rows[index]
        if not (is_name(row[2]) and re.fullmatch(r"1d(?:6|8|10|12)", as_text(row[3]))):
            continue
        add_entry(cursor, source_id, "class", row[2].strip(), index + 1,
                  "图表", "E", {"生命骰": row[3], "首级生命值基础": row[4]})
        counts["class"] += 1

    # U:Z is a weapon catalogue, not the character's equipped weapon rows.
    for index, row in enumerate(rows, 1):
        name, damage = row[20], row[21]
        if not is_name(name) or name.startswith("自定义武器") or not re.fullmatch(r"\d+d\d+(?:/\d+d\d+)?", as_text(damage)):
            continue
        add_entry(cursor, source_id, "weapon", name.strip(), index,
                  "图表", "Z", {"伤害骰": damage, "伤害类型": row[22],
                                    "价格": row[23], "特性": row[24], "重量": row[25]})
        counts["weapon"] += 1

    # AB:AE has adventuring gear, with labels and editable slots interspersed.
    for index, row in enumerate(rows, 1):
        name = row[27]
        if not is_name(name) or name in {"物品名称", "物品描述", "当前名称"}:
            continue
        if as_text(row[29]) in ("", "-") and as_text(row[28]) in ("", "-"):
            continue
        add_entry(cursor, source_id, "item", name.strip(), index, "图表", "AE",
                  {"描述": row[28], "价格": row[29], "重量": row[30]})
        counts["item"] += 1

    # Four parallel background blocks occupy rows 2:63, 66:127, 130:191,
    # 194:255. Keeping the offsets is essential: values are not all on one row.
    for index, row in enumerate(rows[1:], 2):
        name = row[34]
        if not is_name(name) or not is_name(row[35]):
            continue
        if not 2 <= index <= 63:
            continue
        def choices(offset):
            return [as_text(value) for value in rows[index - 1 + offset][40:48]
                    if is_name(value)]
        add_entry(cursor, source_id, "background", name.strip(), index,
                  "图表", "AZ", {"背景特性": row[35], "特殊选项": row[36],
                                  "性格特点": choices(0), "理想": choices(64),
                                  "牵绊": choices(128), "缺点": choices(192),
                                  "技能熟练": row[51],
                                  "工具熟练": rows[index - 1 + 64][51],
                                  "语言": rows[index - 1 + 128][51],
                                  "初始装备": rows[index - 1 + 192][51],
                                  "字段来源": {"性格特点": f"AO{index}:AV{index}",
                                               "理想": f"AO{index+64}:AV{index+64}",
                                               "牵绊": f"AO{index+128}:AV{index+128}",
                                               "缺点": f"AO{index+192}:AV{index+192}",
                                               "技能熟练": f"AZ{index}",
                                               "工具熟练": f"AZ{index+64}",
                                               "语言": f"AZ{index+128}",
                                               "初始装备": f"AZ{index+192}"}})
        counts["background"] += 1

    # CY:CZ is the feat name/description table, separate from the gear list.
    for index, row in enumerate(rows[1:], 2):
        if is_name(row[102]) and isinstance(row[103], str) and len(row[103].strip()) >= 8:
            add_entry(cursor, source_id, "feat", row[102].strip(), index,
                      "图表", "CZ", {"描述": row[103]})
            counts["feat"] += 1

    # F74:G93 contains the 1–20 level spell-point variant lookup. F94 is a
    # character-specific formula row and must not become a rule-table entry.
    for index in range(74, 94):
        level, points = rows[index - 1][5:7]
        if not (isinstance(level, int) and 1 <= level <= 20 and isinstance(points, int)):
            continue
        add_cell_entry(cursor, source_id, "spell_progression", f"法术值变体·{level}级",
                       "图表", f"F{index}:G{index}", {
                           "模式": "法术值变体", "等级": level, "法术值": points,
                           "说明": "仅是原表的法术值查表，不等于法术位自动计算。",
                       })
        counts["spell_progression"] += 1

    # BQ:BW is a race/subrace choice grid, not a complete mechanical statblock.
    for index, row in enumerate(rows[58:], 59):
        family = as_text(row[column_index_from_string("BQ") - 1])
        seen = set()
        for col in range(column_index_from_string("BQ"), column_index_from_string("BW") + 1):
            if col == column_index_from_string("BQ") and family in row[column_index_from_string("BR")-1:column_index_from_string("BW")]:
                continue  # retain the existing subgrid ID when the base name is repeated there
            name = row[col - 1]
            if not is_name(name) or name.strip() in seen:
                continue
            seen.add(name.strip())
            address = f"{get_column_letter(col)}{index}"
            add_cell_entry(cursor, source_id, "species", name, "图表", address,
                           {"族系": family, "说明": "原表种族/亚种选项；不包含完整特性。"})
            counts["species"] += 1

    # BX:CW is 13 pairs of class-feature names and descriptions. Levels are not
    # explicitly stored alongside these cells, so we expose text, not auto-grants.
    first = column_index_from_string("BX")
    last = column_index_from_string("CW")
    for name_col in range(first, last + 1, 2):
        class_name = as_text(rows[0][name_col - 1])
        for index, row in enumerate(rows[1:30], 2):
            name, description = row[name_col - 1], row[name_col]
            if not is_name(name) or not isinstance(description, str) or len(description.strip()) < 8:
                continue
            cell_range = f"{get_column_letter(name_col)}{index}:{get_column_letter(name_col + 1)}{index}"
            add_cell_entry(cursor, source_id, "class_feature", name, "图表", cell_range,
                           {"职业": class_name, "描述": description})
            counts["class_feature"] += 1

    # BF:BG is a separate catalog of selectable maneuvers, metamagic, invocations,
    # infusions, arcane shots, runes and elemental disciplines.
    option_name_col = column_index_from_string("BF") - 1
    category = ""
    for index, row in enumerate(rows, 1):
        name, description = row[option_name_col], row[option_name_col + 1]
        if not is_name(name):
            continue
        if not isinstance(description, str) or len(description.strip()) < 8:
            category = name.strip()
            continue
        add_cell_entry(cursor, source_id, "feature_option", name, "图表",
                       f"BF{index}:BG{index}", {"分类": category, "描述": description})
        counts["feature_option"] += 1
    return counts


def import_workbook(source: Path, destination: Path):
    source = source.resolve(strict=True)
    if source.suffix.lower() not in {".xlsm", ".xlsx"}:
        raise ValueError("只接受 .xlsm / .xlsx 工作簿")
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    source_id = f"5e-workbook-{digest[:16]}"
    destination.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(destination)
    connection.execute("PRAGMA foreign_keys = ON")
    connection.executescript((ROOT / "schemas" / "content-library.sql").read_text(encoding="utf-8"))
    existing = connection.execute("SELECT id FROM library_sources WHERE sha256 = ?", (digest,)).fetchone()
    if existing:
        before_changes = connection.total_changes
        warnings.simplefilter("ignore", UserWarning)
        cached = load_workbook(source, read_only=True, data_only=True, keep_links=False, keep_vba=False)
        try:
            with connection:
                import_chart_entries(connection, existing[0], cached["图表"])
        finally:
            cached.close()
        counts = dict(connection.execute("SELECT kind, COUNT(*) FROM library_entries WHERE source_id = ? GROUP BY kind", (existing[0],)))
        cells = connection.execute("SELECT COUNT(*) FROM library_cells WHERE source_id = ?", (existing[0],)).fetchone()[0]
        changed = connection.total_changes > before_changes
        connection.close()
        return {"status": "updated" if changed else "already_imported", "source_id": existing[0], "cells": cells, "entries": counts}

    warnings.simplefilter("ignore", UserWarning)
    cached = load_workbook(source, read_only=True, data_only=True, keep_links=False, keep_vba=False)
    formulas = load_workbook(source, read_only=True, data_only=False, keep_links=False, keep_vba=False)
    counts = Counter()
    cell_count = 0
    try:
        with connection:
            connection.execute(
                "INSERT INTO library_sources VALUES (?, ?, ?, ?, ?, ?, ?)",
                (source_id, source.name, digest, "dnd5e-2014",
                 datetime.now(timezone.utc).isoformat(), "local_only",
                 "用户提供的模板；含多来源内容，未核实再分发授权；条目未经规则复核。"),
            )
            for sheet in cached.worksheets:
                role = SHEET_ROLES.get(sheet.title, "unclassified")
                connection.execute(
                    "INSERT INTO library_sheets VALUES (?, ?, ?, ?, ?, ?)",
                    (source_id, sheet.title, role, sheet.sheet_state, sheet.max_row, sheet.max_column),
                )
                formula_rows = formulas[sheet.title].iter_rows(values_only=False)
                batch = []
                for value_row, formula_row in zip(sheet.iter_rows(values_only=False), formula_rows):
                    for value_cell, formula_cell in zip(value_row, formula_row):
                        value = value_cell.value
                        formula = formula_cell.value if formula_cell.data_type == "f" else None
                        if formula is not None and not isinstance(formula, str):
                            formula = json_value({"text": getattr(formula, "text", None),
                                                  "ref": getattr(formula, "ref", None)})
                        if value is None and formula is None:
                            continue
                        batch.append((source_id, sheet.title, value_cell.coordinate,
                                      json_value(value) if value is not None else None, formula))
                        if len(batch) >= 1000:
                            connection.executemany("INSERT INTO library_cells VALUES (?, ?, ?, ?, ?)", batch)
                            cell_count += len(batch)
                            batch.clear()
                if batch:
                    connection.executemany("INSERT INTO library_cells VALUES (?, ?, ?, ?, ?)", batch)
                    cell_count += len(batch)
                if sheet.title == "法术大全":
                    counts["spell"] += import_spells(connection, source_id, sheet)
                elif sheet.title == "图表":
                    counts.update(import_chart_entries(connection, source_id, sheet))
    finally:
        cached.close()
        formulas.close()
        connection.close()
    return {"status": "imported", "source_id": source_id, "cells": cell_count,
            "sheets": len(SHEET_ROLES), "entries": dict(counts), "database": str(destination)}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="5E 原始工作簿（只读）")
    parser.add_argument("--database", type=Path, default=DEFAULT_DB)
    args = parser.parse_args()
    print(json.dumps(import_workbook(args.source, args.database), ensure_ascii=False))
