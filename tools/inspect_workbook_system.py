"""Read-only structural inventory of a supplied character workbook.

This deliberately does not calculate formulas, execute macros, or save source files.
"""

from __future__ import annotations

import argparse
import json
import re
import warnings
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


FUNCTION = re.compile(r"(?<![A-Z0-9_.])(?:_xlfn\.)?([A-Z][A-Z0-9_.]*)\s*\(", re.I)
SHEET_REF = re.compile(r"(?:'((?:[^']|'')+)'|([\w\u4e00-\u9fff&]+))!")


def display_value(value) -> str:
    if hasattr(value, "text") and hasattr(value, "ref"):
        return f"array[{value.ref}] {value.text or ''}"
    return str(value)


def inspect(path: Path) -> dict:
    warnings.simplefilter("ignore", UserWarning)
    book = load_workbook(path, read_only=False, data_only=False, keep_links=False,
                         keep_vba=False)
    result = {"source": str(path), "sheets": [], "definedNames": []}
    for name, definition in book.defined_names.items():
        result["definedNames"].append({"name": name, "ref": str(definition.attr_text)[:240]})
    for sheet in book.worksheets:
        functions = Counter()
        cross = Counter()
        first_rows = defaultdict(list)
        formula_examples = []
        data_columns = Counter()
        column_samples = defaultdict(list)
        populated = formulas = 0
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                populated += 1
                data_columns[get_column_letter(cell.column)] += 1
                col = get_column_letter(cell.column)
                if cell.data_type != "f" and len(column_samples[col]) < 5:
                    column_samples[col].append([cell.coordinate, str(value)[:110]])
                if cell.row <= 8 and len(first_rows[cell.row]) < 80:
                    first_rows[cell.row].append([cell.coordinate, display_value(value)[:90]])
                if cell.data_type == "f":
                    formulas += 1
                    formula = display_value(value)
                    functions.update(match.upper() for match in FUNCTION.findall(formula))
                    cross.update((match[0] or match[1]).replace("''", "'")
                                 for match in SHEET_REF.findall(formula))
                    if len(formula_examples) < 18 and (cell.row < 30 or formulas % 100 == 0):
                        formula_examples.append([cell.coordinate, formula[:280]])
        validations = []
        for rule in sheet.data_validations.dataValidation:
            validations.append({"type": rule.type, "ranges": str(rule.sqref)[:300],
                                "formula1": str(rule.formula1 or "")[:240],
                                "formula2": str(rule.formula2 or "")[:120]})
        result["sheets"].append({
            "name": sheet.title, "state": sheet.sheet_state,
            "dimensions": [sheet.max_row, sheet.max_column],
            "populated": populated, "formulas": formulas,
            "functions": functions.most_common(), "crossSheetRefs": cross.most_common(),
            "dataColumns": data_columns.most_common(),
            "columnSamples": dict(column_samples),
            "topRows": dict(first_rows), "formulaExamples": formula_examples,
            "validations": validations,
        })
    book.close()
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--sheet", help="仅输出指定工作表")
    parser.add_argument("--range", dest="cell_range", help="配合 --sheet 输出矩形单元格，包括公式")
    parser.add_argument("--brief", action="store_true")
    args = parser.parse_args()
    if args.cell_range:
        if not args.sheet:
            parser.error("--range 需要 --sheet")
        book = load_workbook(args.source, read_only=True, data_only=False,
                             keep_links=False, keep_vba=False)
        try:
            sheet = book[args.sheet]
            cells = [[f"{cell.coordinate}={display_value(cell.value)[:250]}" for cell in row
                      if cell.value is not None] for row in sheet[args.cell_range]]
            print("\n".join(" | ".join(row) for row in cells if row))
        finally:
            book.close()
        raise SystemExit
    output = inspect(args.source)
    if args.sheet:
        output["sheets"] = [sheet for sheet in output["sheets"] if sheet["name"] == args.sheet]
    if args.brief:
        for sheet in output["sheets"]:
            for key in ("topRows", "formulaExamples", "validations", "dataColumns", "columnSamples"):
                sheet.pop(key)
        output.pop("definedNames")
    print(json.dumps(output, ensure_ascii=False, indent=2))
